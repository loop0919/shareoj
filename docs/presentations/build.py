"""Build the ShareOJ introduction as editable PowerPoint and vector PDF.

Requires python-pptx==1.0.2, reportlab==5.0.1 and IPAexGothic (see README).
"""
from pathlib import Path
import math

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
W, H = 960, 540
INK, GREEN, MINT = '192431', '18794E', '57E3A0'
MUTED, LINE, PAPER, SOFT = '65717C', 'DCE3DF', 'FFFFFF', 'EFF6F1'
FONT = 'IPAexGothic'
font_path = Path('/usr/share/fonts/opentype/ipaexfont-gothic/ipaexg.ttf')
if not font_path.exists():
    raise SystemExit('Install IPAexGothic; see docs/presentations/README.md')
pdfmetrics.registerFont(TTFont(FONT, str(font_path)))
prs = Presentation()
prs.slide_width, prs.slide_height = Pt(W), Pt(H)
prs.core_properties.title = 'ShareOJ | プロダクト概要と技術構成'
prs.core_properties.subject = 'エンジニア向け紹介 / 2026-09-22 / repository c59a2e1'
prs.core_properties.author = 'ShareOJ'
pdf = canvas.Canvas(str(OUT / 'shareoj-introduction.pdf'), pagesize=(W, H))
pdf.setTitle(prs.core_properties.title)
pdf.setAuthor('ShareOJ')
notes = []
layout_errors = []
slide = None
dark = False


def color(c):
    return tuple(int(c[i:i+2], 16) / 255 for i in (0, 2, 4))


def rect(x, y, w, h, fill=SOFT, stroke=None, ellipse=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL if ellipse else MSO_SHAPE.RECTANGLE,
                                   Pt(x), Pt(y), Pt(w), Pt(h))
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        pdf.setFillColorRGB(*color(fill))
    else:
        shape.fill.background()
    if stroke:
        shape.line.color.rgb = RGBColor.from_string(stroke)
        shape.line.width = Pt(1)
        pdf.setStrokeColorRGB(*color(stroke))
    else:
        shape.line.fill.background()
    if ellipse:
        pdf.ellipse(x, H-y-h, x+w, H-y, fill=bool(fill), stroke=bool(stroke))
    else:
        pdf.rect(x, H-y-h, w, h, fill=bool(fill), stroke=bool(stroke))


def line(x1, y1, x2, y2, c=LINE, width=1, arrow=False):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    shape.line.color.rgb = RGBColor.from_string(c)
    shape.line.width = Pt(width)
    pdf.setStrokeColorRGB(*color(c))
    pdf.setLineWidth(width)
    pdf.line(x1, H-y1, x2, H-y2)
    if arrow:
        angle = math.atan2(y2-y1, x2-x1)
        for offset in (-.5, .5):
            line(x2, y2, x2-8*math.cos(angle+offset), y2-8*math.sin(angle+offset), c, width)


def text(s, x, y, w=850, size=20, c=None, bold=False, leading=1.4):
    c = c or (PAPER if dark else INK)
    lines = s.split('\n')
    height = len(lines)*size*leading + 4
    assert x >= 0 and y >= 0 and x+w <= W+.1 and y+height <= H, (s, x, y, w, height)
    for part in lines:
        if pdfmetrics.stringWidth(part, FONT, size) > w:
            layout_errors.append((len(prs.slides), part, w, round(pdfmetrics.stringWidth(part, FONT, size))))
    box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, part in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = part
        p.font.name, p.font.size, p.font.bold = FONT, Pt(size), bold
        p.font.color.rgb = RGBColor.from_string(c)
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = Pt(size*leading)
        for run in p.runs:
            ea = OxmlElement('a:ea')
            ea.set('typeface', FONT)
            run._r.get_or_add_rPr().append(ea)
        pdf.setFillColorRGB(*color(c))
        pdf.setFont(FONT, size)
        pdf.drawString(x, H-y-size*.95-i*size*leading, part)
    return box


def page(title, section, note, sources=(), inverted=False):
    global slide, dark
    if slide is not None:
        pdf.showPage()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    dark = inverted
    rect(0, 0, W, H, INK if dark else PAPER)
    text(section, 48, 25, size=11, c=MINT if dark else GREEN)
    if title:
        text(title, 48, 65, size=30, bold=True)
    line(48, 493, 912, 493, '3D4952' if dark else LINE)
    n = len(prs.slides)
    text('ShareOJ  /  Product & Engineering', 48, 507, 560, 9, MUTED if not dark else 'B7C7BF')
    text(f'{n:02d} / 17', 845, 505, 70, 11, MUTED if not dark else 'B7C7BF')
    for source in sources:
        assert (ROOT / source).exists(), source
    full_note = note + '\n\n根拠（基準コミット c59a2e1）:\n' + '\n'.join(sources)
    slide.notes_slide.notes_text_frame.text = full_note
    notes.append(f'## {n:02d} {title or "ShareOJ"}\n\n{note}\n\n' + '\n'.join(f'- [{p}](../../{p})' for p in sources))


def node(title, sub, x, y, w=176, h=78, fill=SOFT):
    rect(x, y, w, h, fill)
    text(title, x+14, y+12, w-28, 18, GREEN, True)
    if sub:
        text(sub, x+14, y+43, w-28, 12, INK)


page('', 'SHARE ONLINE JUDGE',
     '約15分でShareOJの利用体験と採点基盤を紹介する。2026年9月22日のリポジトリを根拠とし、本番環境をこの資料作成時に再検証したものではない。',
     ['README.md', 'web/app/pages/index.vue'], True)
text('ShareOJ', 48, 109, 610, 64, bold=True)
text('考える楽しさを、\n次の一問へ。', 48, 214, 650, 38, leading=1.5)
text('プロダクト概要と技術構成', 50, 367, 630, 22, 'C8D8D0')
text('ENGINEERING INTRODUCTION  /  2026.09.22', 50, 433, 760, 11, MINT)
for a, b in [((682,320),(802,239)),((802,239),(864,135)),((682,320),(714,167)),((714,167),(864,135))]:
    line(*a,*b,'3B554A',2)
line(682,320,802,239,MINT,3)
line(802,239,864,135,MINT,3)
for x,y,r in [(682,320,18),(802,239,9),(864,135,16),(714,167,7)]:
    rect(x-r,y-r,r*2,r*2,MINT if x!=714 else '3B554A',ellipse=True)

page('解く経験を、次の問題と知識へ', '01  PRODUCT',
     'ShareOJはプログラミング問題を解き、作り、共有できるオンラインジャッジ。利用者は解答者だけでなく作問者やコンテスト主催者にもなれる。4つの機能を学習と創作のつながりとして紹介する。', ['README.md'])
for i,(title,body) in enumerate([
    ('解く','コードを提出し、\n結果から解法を見直す。'),
    ('作る','問題文とテストを用意し、\n自作の問題を公開する。'),
    ('共有する','解き方や学んだことを\n記事として残す。'),
    ('集まる','コンテストを開催し、\n同じ問題に挑戦する。')]):
    x=48+i*222
    text(f'0{i+1}',x,161,180,42,GREEN)
    line(x,225,x+195,225)
    text(title,x,247,196,27,bold=True)
    text(body,x,300,205,16,MUTED,leading=1.7)
text('解答者から作問者へ。ひとつのサービスで活動を広げられる。',48,424,860,22)

page('問題を読んで、提出し、結果を確かめる', '02  SOLVING',
     '画面そのもののスクリーンショットではなく、利用の流れを示す模式図。サンプル検証と本提出は別操作。進捗は実際の採点情報から表示し、結果詳細にはケースの判定や資源使用量がある。ACなどの表示は説明用の例で、実測や本番提出の結果ではない。',
     ['web/app/components/SubmissionForm.vue','web/app/components/SubmissionStatus.vue','web/app/content/guides/language-guide.md'])
for i,t in enumerate(['問題を読む','サンプル検証','提出する','結果を確認']):
    node(t,'',48+i*222,139,198,52)
    if i<3: line(250+i*222,165,265+i*222,165,GREEN,1.5,True)
rect(48,220,400,215,INK)
text('CODE  /  Python の例',68,239,350,12,MINT)
text('a, b = map(int, input().split())\nprint(a + b)',68,283,357,16,PAPER,leading=1.8)
text('標準入力 → プログラム → 標準出力',68,389,358,13,'B7C7BF')
text('RESULT  /  表示内容の例',490,234,400,12,GREEN)
for i,(v,t) in enumerate([('AC','正解'),('WA','出力が条件を満たさない'),('TLE','実行時間の制限を超過')]):
    y=275+i*53
    text(v,490,y,62,22,GREEN if i==0 else INK,True)
    text(t,573,y+3,333,17)
    line(490,y+40,912,y+40)

page('作問を支える編集と検証', '03  AUTHORING',
     '問題編集画面には問題文、解説、テストケース、生成、判定コード、管理の入口がある。Markdownと数式、生成コードと入力検証、未公開問題への提出、テスター招待を組み合わせて公開前に確認できる。古いweb/READMEの未実装という記述より現行コンポーネントを優先した。',
     ['web/app/pages/problems/new.vue','web/app/content/guides/generator-guide.md','api/README.md'])
text('書く',48,153,180,36,GREEN)
text('Markdownと数式で\n問題文と解説を編集',250,157,630,22)
line(48,234,912,234)
text('確かめる',48,257,200,36,GREEN)
text('テストケース、生成コード、入力検証を用意\n未公開のまま提出し、テスターと確認',250,256,660,21)
line(48,338,912,338)
text('公開する',48,361,200,36,GREEN)
text('問題を公開、またはコンテストへ登録\n提出時の問題と採点条件を記録',250,360,660,21)

page('コンテストと記事で、学びを共有する', '04  COMMUNITY',
     'コンテストは主催者が開催期間、問題、配点、誤答ペナルティを設定する。公式順位は期間内にコンテストの問題ページから受け付けた提出が対象。記事は問題の解き方や学びを共有する場である。利用者数、継続率、開催実績などの定量効果は確認していないため記載しない。',
     ['web/app/pages/blog/contest-rules.vue','web/app/pages/blog/new.vue','README.md'])
text('CONTEST',48,149,390,13,GREEN)
text('同じ問題に挑む場',48,183,405,29,bold=True)
text('開催期間、問題、配点を設定\n得点とペナルティで順位を表示\n終了後は問題と登録済み解説を公開',48,253,414,19,leading=1.9)
line(472,152,472,420)
text('ARTICLE',518,149,390,13,GREEN)
text('考え方を残す場',518,183,394,29,bold=True)
text('解き方や学んだことを記事にする\n問題に取り組む人が読める\n解答の先にある理解を共有する',518,253,394,19,leading=1.9)

page('Webと採点基盤を分けた全体構成', '05  ARCHITECTURE',
     'Nuxt SSRとGo APIはそれぞれAPI Gateway HTTP APIとLambdaで公開。図ではHTTP APIをノード内の説明へまとめた。Go APIとbridgeはPostgreSQLに接続し、workerはSQSとS3を介して動く。線は主な要求またはデータ移動の向きを示す。結果はbridgeからDBへ反映され、画面側がAPI経由で取得する。',
     ['infra/README.md','docs/judge/execution-model.md','judge/README.md'])
for x,title,sub in [(48,'Browser','閲覧 / 編集 / 提出'),(270,'Nuxt 4 + TS','SSR / HTTP API + Lambda'),(492,'Go API','HTTP API + Lambda'),(714,'PostgreSQL','RDS / 提出とOutbox')]:
    node(title,sub,x,142,198)
for x in [246,468,690]: line(x+2,181,x+20,181,GREEN,1.5,True)
node('Cognito','ログイン / トークン検証',270,264,198)
line(370,220,370,264,GREEN,1.5,True)
node('Bridge Lambda','配送 / 進捗と結果の反映',714,264,198)
line(813,264,813,220,GREEN,1.5,True)
line(826,220,826,264,GREEN,1.5,True)
node('Lightsail × 2','worker / isolate',48,382,198)
node('SQS','要求 / 進捗と結果',370,382,252)
node('S3','ソース / テストデータ',714,382,198)
line(714,303,652,303,GREEN,1.5)
line(652,303,652,366,GREEN,1.5)
line(652,366,496,366,GREEN,1.5)
line(496,366,496,382,GREEN,1.5,True)
line(652,303,714,303,GREEN,1.5,True)
line(370,411,246,411,GREEN,1.5,True)
line(246,435,370,435,GREEN,1.5,True)
line(813,342,813,382,GREEN,1.5,True)
line(714,470,147,470,GREEN,1.5)
line(813,460,813,470,GREEN,1.5)
line(813,470,714,470,GREEN,1.5)
line(147,470,147,460,GREEN,1.5,True)

page('1件の提出が結果になるまで', '06  JUDGING PIPELINE',
     '提出行が永続Outboxを兼ねる。bridgeは未配送の提出を読み、入力をS3に保存して要求をSQSへ送る。workerは入力とruntime digestを検証し、コンパイルが必要な言語では一度コンパイルした成果物をケース間で使う。進捗と最終結果を結果キューで返し、bridgeがDBへ反映する。UIのポーリングは現行SubmissionFormで1.5秒間隔。',
     ['api/internal/database/005_judge_outbox.sql','api/cmd/judge-bridge/dispatch.go','api/cmd/judge-bridge/results.go','web/app/components/SubmissionForm.vue'])
for i,(name,body) in enumerate([
    ('受付','問題の版、言語、\n制限値をDBに保存'),('配送','S3に入力を置き、\nSQSへ要求を送る'),
    ('採点','入力を検証し、\n隔離環境で実行'),('確定','結果をDBへ反映し、\n画面から取得')]):
    x=48+222*i
    text(f'{i+1:02}',x,149,195,46,GREEN)
    text(name,x,222,195,28,bold=True)
    text(body,x,275,203,18,leading=1.7)
    if i<3: line(x+159,174,x+205,174,GREEN,2,True)
rect(48,382,864,69,SOFT)
text('受付と採点を非同期に分離し、処理中の進捗も返す。',69,401,824,24,GREEN)

page('再配信があっても、確定結果を守る', '07  DELIVERY & CONSISTENCY',
     'SQS Standardの再配信を前提にしている。exactly-once実行を保証するという意味ではない。配送時にはFOR UPDATE SKIP LOCKEDを使用。結果は提出IDと試行IDを照合し、確定済み提出は上書きしない。進捗の逆行もSQL条件で抑える。worker停止時は可視性期限後の再配信を使い、期限を過ぎた未完了提出はdispatch時にJEへ確定する。',
     ['api/cmd/judge-bridge/dispatch.go','api/cmd/judge-bridge/results.go','api/internal/submissions/finish.go','docs/judge/execution-model.md'])
text('起こり得ること',48,146,320,14,MUTED)
text('実装での扱い',387,146,525,14,MUTED)
for y,a,b in [(194,'API受付後の配送待ち','提出行をOutboxとしてDBに残す'),(270,'要求や結果の重複','試行IDを照合し、確定結果を上書きしない'),(346,'ワーカーの途中停止','SQSが再配信し、未完了のまま残さない')]:
    line(48,y-14,912,y-14)
    text(a,48,y,315,21)
    text(b,387,y,525,20,GREEN)
text('「一度だけ実行」ではなく、重複を受け止める設計。',48,436,865,21)

page('提出コードをケースごとに隔離する', '08  SANDBOX',
     '通常判定の模式図。isolateは名前空間、UID、cgroup v2で提出コードを隔離する。各ケースの作業領域と子孫プロセスを回収し、コンパイル成果物だけを次ケースへ渡す。期待出力や認証情報は解答プログラム側へ渡さない。スペシャルジャッジと対話用ジャッジには、それぞれ別の隔離環境へ必要なテストデータを渡す。共有カーネルなのでVM相当の隔離とは主張しない。',
     ['judge/README.md','docs/judge/execution-model.md','docs/judge/interactive-judge.md'])
rect(48,143,550,291,SOFT,LINE)
text('Lightsail host',67,158,500,15,GREEN)
text('管理ワーカー',67,192,500,24,bold=True)
rect(80,249,483,153,PAPER,GREEN)
text('isolate / ケース専用cgroup',99,264,441,20,GREEN)
text('実行ファイル ＋ 入力\nCPU、メモリ、出力を制限\n外部ネットワークへの経路なし',99,307,438,18,leading=1.5)
text('提出側へ渡さない',644,162,267,23,bold=True)
text('AWS認証情報\n管理ワーカーのコード\n期待出力と計測メタデータ',644,215,269,17,leading=1.9)
text('ケース終了後に\n作業領域と子孫を回収',644,347,269,21,GREEN)
text('境界：提出と管理処理は同じカーネルを共有する。',48,455,864,17,MUTED)

page('資源を制限し、判定の根拠を記録する', '09  RESOURCE ACCOUNTING',
     '通常提出の問題設定はCPU時間100〜5000ms、メモリ64〜512MiB。メモリは子孫と課金対象のキャッシュを含むcgroupピークで、RSSとは区別する。通常判定の経過時間上限はCPU上限の3倍+1秒。対話形式には別の経過時間計算がある。既存2台構成は1台1提出で最大2提出。同一ホスト内でコンパイルと通常提出のケース実行を重ねない。',
     ['docs/judge/runtime-policy.md','judge/README.md','docs/judge/interactive-judge.md'])
for x,num,unit,label in [(48,'100–5,000','ms','ケースのCPU時間'),(354,'64–512','MiB','問題ごとのメモリ上限'),(660,'1','提出 / ホスト','ホスト内の同時採点')]:
    text(num,x,158,275,38,GREEN)
    text(unit,x,217,275,18,MUTED)
    text(label,x,261,279,19)
line(48,314,912,314)
text('TLE',48,343,85,25,GREEN)
text('CPU時間と経過時間を別々に監視する。',169,346,741,21)
text('MLE',48,397,85,25,GREEN)
text('ケースのcgroup OOMを根拠にする。\nホスト障害や計測不能はJEとして区別する。',169,397,741,19,leading=1.5)

page('問題に合わせて選ぶ3つの判定方法', '10  JUDGE MODES',
     '通常判定は期待出力との比較、スペシャルジャッジは作問者の検証コード、対話形式は提出と対話用ジャッジの双方向通信。チェッカーとインタラクターは別の隔離環境で実行し、資源を制限する。対話では両側の正常終了を待ってACにする。スコアファイルの値の集計や部分点を提供するという意味ではない。',
     ['docs/judge/special-judge.md','docs/judge/interactive-judge.md'])
for y,name,body,flow in [(148,'通常判定','出力を期待出力と比較','入力 → 提出 → 出力比較'),(251,'スペシャルジャッジ','複数の正解を検証コードで判定','提出 → 出力 → 検証コード'),(354,'対話形式','質問と応答を交わす問題','提出 ↔ 中継 ↔ 対話用ジャッジ')]:
    line(48,y-12,912,y-12)
    text(name,48,y+3,284,24,GREEN)
    text(body,355,y,557,19)
    text(flow,355,y+41,557,20,MUTED)

page('言語環境を固定し、検証して公開する', '11  RUNTIMES',
     '言語一覧はリポジトリの言語ガイドに基づく例。受付可能なランタイムはGET /runtimesと提出欄が正とし、ここでは本番の公開状態を再確認していない。各言語の配布元、版、ライブラリをlockファイルで固定し、runtime名とdigestで採点条件を識別する。旧環境を保存して過去提出を再実行する仕組みは未実装。',
     ['web/app/content/guides/language-guide.md','docs/judge/runtime-policy.md','judge/runtimes.py'])
text('言語ガイドに掲載される環境',48,148,460,15,MUTED)
text('C / C++ / Rust\nPython / Java / C#\nGo / Nim / Haskell\nJavaScript / TypeScript / Ruby',48,189,495,25,leading=1.75)
line(570,150,570,419)
for i,(a,b) in enumerate([('固定','コンパイラとライブラリの版'),('検証','実機smokeと資源制限'),('公開','digestと受付対象を同期')]):
    y=164+88*i
    text(a,615,y,95,26,GREEN)
    text(b,615,y+41,296,16)
text('実際に選べる言語は、提出欄と GET /runtimes で確認。',48,449,864,18,MUTED)

page('SSRと認証をサーバー側でつなぐ', '12  WEB & AUTH',
     'Nuxtは問題文や数式を含むHTMLをSSRする。トークンはHttpOnly Cookieに保存し、NuxtサーバーからGo APIへBearerとして渡す。Cookieを使う書き込みにはOrigin検証を適用。Markdownの生HTMLや危険なURLを抑え、KaTeXはtrust:falseで描画する。認証済みであることに加え、Go APIは所有者や公開範囲などの認可を行う。',
     ['web/README.md','api/README.md','web/app/utils/problem-markdown.ts'])
node('Browser','HttpOnly Cookie',48,153,222)
node('Nuxt server','SSR / セッション更新',365,153,222)
node('Go API','Bearer検証 / 認可',682,153,230)
line(270,192,363,192,GREEN,2,True)
line(587,192,680,192,GREEN,2,True)
for y,a,b in [(277,'読み始められる','初回HTMLに問題文と数式を含める'),(335,'権限を確かめる','認証に加えて、所有者と公開範囲を検証'),(393,'入力を制御する','Markdownと数式の安全な描画、Origin検証')]:
    text(a,48,y,274,23,GREEN)
    text(b,365,y+3,547,19)

page('配布の完了を、提出の成功まで確認する', '13  DELIVERY & OPERATIONS',
     'WebとAPIはGitHub Actionsで型、ビルド、ブラウザー、実DBなどを検証する。ジャッジ更新はrollout.py runに統合され、受付停止、DBとキューの排出、配布、全言語検証、起動と設定同期、受付再開、本番提出の検証まで行う。SSMへ登録しただけでは完了としない。今回の作業では配布や本番提出は実施していない。',
     ['.github/workflows/ci.yml','docs/judge/deployment-checks.md'])
text('WEB / API',48,145,850,12,GREEN)
text('型検査 → ビルド → ブラウザー検証 → 実DB連携テスト',48,177,864,22)
line(48,232,912,232)
text('JUDGE',48,257,850,12,GREEN)
for i,t in enumerate(['受付停止','DBとキュー\nを排出','配布と\n実機検証','設定同期と\n受付再開','本番提出\nを検証']):
    x=48+176*i
    rect(x,298,160,86,SOFT)
    text(t,x+12,314,138,19,GREEN,leading=1.5)
    if i<4: line(x+161,341,x+174,341,GREEN,1.5,True)
text('完了条件：state.json の status: passed / step: complete',48,425,864,21)

page('設計上の選択と、現在の制約', '14  TRADE-OFFS',
     '左の設計選択に対して、右に運用上の制約を対応させる。2台は既存配布対象としてドキュメントに記録されている構成であり、今回稼働状態を検査したものではない。RDS Single-AZはリポジトリの環境設定に基づく。性能やセキュリティの保証、無停止配布、無制限のスケーラビリティは主張しない。',
     ['judge/README.md','infra/README.md','docs/judge/runtime-policy.md','docs/judge/deployment-checks.md'])
text('採用した構成',48,143,335,14,MUTED)
text('運用で考慮すること',414,143,498,14,MUTED)
for y,a,b in [(192,'Lightsail + isolate','共有カーネル。CPU負荷で計測値が変動し得る。'),(263,'2台で要求キューを共有','最大2提出を同時採点。更新時は受付を停止。'),(334,'RDS PostgreSQL','設定はSingle-AZ。自動フェイルオーバーなし。'),(405,'runtime digestで条件を識別','過去の実行環境を保存する仕組みは未実装。')]:
    line(48,y-12,912,y-12)
    text(a,48,y,351,21,GREEN)
    text(b,414,y+2,498,17)

page('考える楽しさを、次の一問へ。', '15  SHAREOJ',
     '締めは、使う側と技術側の説明を接続する。ShareOJは問題を解くだけでなく、作問、記事、コンテストまでをつなぐ。その下で、Web/API、永続化、非同期配送、隔離実行が役割を分担している。公開サイトとリポジトリへ案内する。',
     ['README.md'],True)
text('解く。作る。共有する。',48,170,864,42,bold=True)
text('問題と記事、コンテストをつなぐオンラインジャッジ。',48,254,864,23,'C8D8D0')
text('Nuxt + Go + PostgreSQL\nS3 / SQS + Lightsail / isolate',48,322,864,24,MINT,leading=1.5)
b=text('www.share-oj.net',48,431,440,22,PAPER)
b.text_frame.paragraphs[0].runs[0].hyperlink.address='https://www.share-oj.net/'
pdf.linkURL('https://www.share-oj.net/',(48,H-465,488,H-431),relative=0)

page('参照資料と、このスライドの前提', 'APPENDIX  /  SOURCES',
     '各スライドの発表者ノートに参照ファイルを記載した。基準コミットはc59a2e1。旧ADRや初期実装の説明と現行コードに差がある箇所は、新しい運用資料と言語ガイド、実装を優先した。数値は設定値であり、今回の負荷試験の実測結果ではない。図は説明用であり、実画面のスクリーンショットではない。',
     ['README.md','docs/judge/runtime-policy.md','docs/judge/deployment-checks.md'])
for y,a,b in [(144,'概要と利用体験','README.md / web/app/pages/'),(199,'WebとAPI','web/README.md / api/README.md / infra/README.md'),(254,'配送と結果反映','api/cmd/judge-bridge/ / api/internal/submissions/'),(309,'採点と隔離','judge/README.md / docs/judge/runtime-policy.md'),(364,'公開と検証','docs/judge/deployment-checks.md / .github/workflows/ci.yml')]:
    text(a,48,y,240,19,GREEN)
    text(b,310,y+2,602,13)
text('基準：2026-09-22 / repository c59a2e1\n本番環境の再検証、利用実績や性能の測定は本資料の対象外。',48,433,864,14,MUTED,leading=1.5)

assert len(prs.slides) == 17
assert not layout_errors, layout_errors
pdf.save()
prs.save(OUT / 'shareoj-introduction.pptx')
(OUT / 'speaker-notes.md').write_text('# ShareOJ 発表者ノート\n\n想定時間は約15分（付録を除く）。\n\n'+'\n\n'.join(notes)+'\n',encoding='utf-8')
print(f'Created {len(prs.slides)} slides: editable PPTX, vector PDF, speaker notes')
